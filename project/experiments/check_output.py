"""
experiments/check_output.py

Reads test_output.json, prints a console summary, and exports an Excel file
with two sheets:
  - Summary   : overall counts
  - Papers     : one row per paper with id, eq count, status, eq numbers, methods

Run from the project root:
    python project/experiments/check_output.py
"""

import json
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

_HERE       = Path(__file__).resolve().parent
_PROJECT    = _HERE.parent
OUTPUT_PATH = _PROJECT / "data" / "output" / "test_output.json"
EXCEL_PATH  = _PROJECT / "data" / "output" / "equations_check.xlsx"

MAX_EQ = 7

# ── colour palette ────────────────────────────────────────────────────────────
GREEN  = "C6EFCE"   # full (7 eqs)
YELLOW = "FFEB9C"   # under 7
RED    = "FFC7CE"   # empty
HEADER = "4472C4"   # header row blue


def _cell_color(ws, row, col, hex_color):
    ws.cell(row=row, column=col).fill = PatternFill(
        fill_type="solid", fgColor=hex_color
    )


def main() -> None:
    with open(OUTPUT_PATH, encoding="utf-8") as f:
        data = json.load(f)

    # ── build rows ────────────────────────────────────────────────────────────
    rows = []
    for arxiv_id, equations in data.items():
        count = len(equations)
        if count == 0:
            status = "EMPTY"
        elif count < MAX_EQ:
            status = "UNDER"
        else:
            status = "FULL"

        eq_nums   = ", ".join(equations.keys())
        methods   = list({d.get("latex_method", "") for d in equations.values()})
        sources   = list({d.get("source", "") for d in equations.values()})

        rows.append({
            "arxiv_id" : arxiv_id,
            "count"    : count,
            "status"   : status,
            "eq_nums"  : eq_nums,
            "methods"  : ", ".join(sorted(methods)),
            "source"   : ", ".join(sorted(sources)),
        })

    full    = [r for r in rows if r["status"] == "FULL"]
    under   = sorted([r for r in rows if r["status"] == "UNDER"], key=lambda r: r["count"])
    empty   = [r for r in rows if r["status"] == "EMPTY"]

    # ── console output ────────────────────────────────────────────────────────
    total_eq = sum(r["count"] for r in rows)
    print(f"Output file    : {OUTPUT_PATH}")
    print(f"Total papers   : {len(rows)}")
    print(f"Total equations: {total_eq}")
    print(f"Full (7 eqs)   : {len(full)}")
    print(f"Under 7 eqs    : {len(under)}")
    print(f"Empty (0 eqs)  : {len(empty)}")
    print()

    if under:
        print("Papers with fewer than 7 equations:")
        print("-" * 55)
        for r in under:
            print(f"  {r['arxiv_id']}  ->  {r['count']} eq(s)  [{r['eq_nums']}]")
        print()

    if empty:
        print("Empty papers (0 equations):")
        print("-" * 55)
        for r in empty:
            print(f"  {r['arxiv_id']}")
        print()

    # ── Excel export ──────────────────────────────────────────────────────────
    wb = openpyxl.Workbook()

    # ── Sheet 1: Summary ──────────────────────────────────────────────────────
    ws_sum = wb.active
    ws_sum.title = "Summary"

    summary_data = [
        ("Metric",          "Value"),
        ("Total papers",    len(rows)),
        ("Total equations", total_eq),
        ("Full (7 eqs)",    len(full)),
        ("Under 7 eqs",     len(under)),
        ("Empty (0 eqs)",   len(empty)),
    ]

    for i, (label, value) in enumerate(summary_data, start=1):
        ws_sum.cell(row=i, column=1, value=label)
        ws_sum.cell(row=i, column=2, value=value)

    # header row styling
    for col in range(1, 3):
        cell = ws_sum.cell(row=1, column=col)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(fill_type="solid", fgColor=HEADER)

    ws_sum.column_dimensions["A"].width = 22
    ws_sum.column_dimensions["B"].width = 14

    # ── Sheet 2: Papers ───────────────────────────────────────────────────────
    ws = wb.create_sheet("Papers")

    headers = ["arXiv ID", "Eq Count", "Status", "Equation Numbers", "LaTeX Method", "Source"]
    for col, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(fill_type="solid", fgColor=HEADER)
        cell.alignment = Alignment(horizontal="center")

    # write rows: full first, then under, then empty
    row_idx = 2
    for group in [full, under, empty]:
        for r in group:
            ws.cell(row=row_idx, column=1, value=r["arxiv_id"])
            ws.cell(row=row_idx, column=2, value=r["count"])
            ws.cell(row=row_idx, column=3, value=r["status"])
            ws.cell(row=row_idx, column=4, value=r["eq_nums"])
            ws.cell(row=row_idx, column=5, value=r["methods"])
            ws.cell(row=row_idx, column=6, value=r["source"])

            color = GREEN if r["status"] == "FULL" else (YELLOW if r["status"] == "UNDER" else RED)
            for col in range(1, 7):
                _cell_color(ws, row_idx, col, color)
                ws.cell(row=row_idx, column=col).alignment = Alignment(horizontal="left")

            row_idx += 1

    # column widths
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 10
    ws.column_dimensions["C"].width = 10
    ws.column_dimensions["D"].width = 45
    ws.column_dimensions["E"].width = 16
    ws.column_dimensions["F"].width = 10

    # freeze header row
    ws.freeze_panes = "A2"

    wb.save(EXCEL_PATH)
    print(f"Excel saved to : {EXCEL_PATH}")


if __name__ == "__main__":
    main()
